
import re
import shutil
import subprocess
import tempfile
import zipfile
from pathlib import Path

import pandas as pd
import streamlit as st
from docx import Document
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT

from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.enums import TA_CENTER
from reportlab.lib.units import mm
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, Image

st.set_page_config(page_title="Hall Ticket Generator", page_icon="🎫", layout="wide")

BASE = Path(__file__).parent
TEMPLATE = BASE / "Hall_Ticket_Format.docx"
HEADER_IMAGE = BASE / "header.png"

STUDENT_COLS = [
    "Student Name", "Seat / Roll No.", "PRN / Enrollment No.",
    "Class", "Course / Programme", "Semester", "Examination Centre"
]
SCHEDULE_COLS = ["Sr. No.", "Date", "Day", "Subject / Paper", "Time"]

def safe_name(value):
    value = re.sub(r'[\\/:*?"<>|]+', "_", str(value).strip())
    return re.sub(r"\s+", " ", value)[:100] or "Student"

def clean_br(text):
    # Remove literal HTML break tags that should never appear in Word/PDF output.
    return re.sub(r"<\s*br\s*/?\s*>", "\n", str(text))

def set_normal_style(doc):
    style = doc.styles["Normal"]
    style.font.name = "Times New Roman"
    style.font.size = Pt(9)

def make_docx(student, schedule, exam_title, photo_path, out_path):
    doc = Document(TEMPLATE)
    set_normal_style(doc)

    # Force a compact one-page layout.
    sec = doc.sections[0]
    sec.top_margin = Inches(0.32)
    sec.bottom_margin = Inches(0.30)
    sec.left_margin = Inches(0.42)
    sec.right_margin = Inches(0.42)

    # Remove any literal <br> tags already present in the template.
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    for run in p.runs:
                        if "<br" in run.text.lower():
                            run.text = clean_br(run.text)

    info = doc.tables[0]
    vals = [
        student.get("Student Name", ""),
        student.get("Seat / Roll No.", ""),
        student.get("PRN / Enrollment No.", ""),
        student.get("Class", ""),
        student.get("Course / Programme", ""),
        student.get("Semester", ""),
        student.get("Examination Centre", "")
    ]
    for i, val in enumerate(vals):
        info.cell(i, 1).text = "" if pd.isna(val) else str(val)

    # Photo placeholder: exactly bracketed wording when no photo is supplied.
    photo_cell = info.cell(0, 2)
    photo_cell.text = ""
    p = photo_cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(0)
    if photo_path and Path(photo_path).exists():
        run = p.add_run()
        run.add_picture(photo_path, width=Inches(1.22))
    else:
        r = p.add_run("(Paste Photograph Here)")
        r.bold = True
        r.font.size = Pt(8)

    # Clean and standardize the schedule header.
    sched = doc.tables[1]
    sched.cell(0, 5).text = "Jr. Supervisor\nSignature"
    sched.cell(0, 5).paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER

    # Keep the supplied 8-row schedule.
    for r in range(1, len(sched.rows)):
        for c in range(5):
            sched.cell(r, c).text = ""
        sched.cell(r, 5).text = ""

    for i, (_, row) in enumerate(schedule.iterrows(), start=1):
        if i >= len(sched.rows):
            break
        for c, col in enumerate(SCHEDULE_COLS):
            sched.cell(i, c).text = "" if pd.isna(row.get(col, "")) else clean_br(row.get(col, ""))

    # Replace exam heading only. Do NOT add an extra academic-year paragraph:
    # that extra paragraph was causing the Word output to spill onto page 2.
    for p in doc.paragraphs:
        if p.text.strip() == "EXAMINATION":
            p.text = exam_title.strip() or "EXAMINATION"
            if p.runs:
                p.runs[0].bold = True
                p.runs[0].font.size = Pt(10)
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            break

    # Compact all paragraphs inside tables.
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                for p in cell.paragraphs:
                    p.paragraph_format.space_before = Pt(0)
                    p.paragraph_format.space_after = Pt(0)
                    p.paragraph_format.line_spacing = 1.0
                    for run in p.runs:
                        if run.font.size is None:
                            run.font.size = Pt(8.5)

    # Compact instructions and footer text.
    for p in doc.paragraphs:
        if p.text.strip().startswith(("1.", "2.", "3.", "4.", "5.")):
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(1)
            p.paragraph_format.line_spacing = 1.0
            for run in p.runs:
                run.font.size = Pt(8)

    doc.save(out_path)

def esc(text):
    return (
        str(text).replace("&", "&amp;")
        .replace("<", "&lt;")
        .replace(">", "&gt;")
    )

def make_pdf(student, schedule, exam_title, photo_path, out_path):
    styles = getSampleStyleSheet()
    normal = ParagraphStyle(
        "normal", parent=styles["Normal"], fontName="Times-Roman",
        fontSize=7.4, leading=8.5
    )
    bold = ParagraphStyle(
        "bold", parent=normal, fontName="Times-Bold"
    )
    center = ParagraphStyle(
        "center", parent=bold, alignment=TA_CENTER, fontSize=7.6, leading=8.5
    )
    title = ParagraphStyle(
        "title", parent=bold, alignment=TA_CENTER, fontSize=13, leading=14
    )
    subtitle = ParagraphStyle(
        "subtitle", parent=bold, alignment=TA_CENTER, fontSize=9.5, leading=10.5
    )

    doc = SimpleDocTemplate(
        str(out_path), pagesize=A4,
        leftMargin=10*mm, rightMargin=10*mm,
        topMargin=7*mm, bottomMargin=6*mm
    )
    story = []

    if HEADER_IMAGE.exists():
        story.append(Image(str(HEADER_IMAGE), width=190*mm, height=36.5*mm))
        story.append(Spacer(1, 1*mm))

    story.append(Paragraph("HALL TICKET", title))
    story.append(Paragraph(esc(exam_title.strip() or "EXAMINATION"), subtitle))
    story.append(Spacer(1, 2*mm))

    labels = [
        "Student Name", "Seat / Roll No.", "PRN / Enrollment No.",
        "Class", "Course / Programme", "Semester", "Examination Centre"
    ]
    data = []
    for lab in labels:
        data.append([Paragraph(esc(lab), bold), Paragraph(esc(student.get(lab, "")), normal), ""])
    if photo_path and Path(photo_path).exists():
        data[0][2] = Image(str(photo_path), width=30*mm, height=35*mm)
    else:
        data[0][2] = Paragraph("(Paste Photograph Here)", center)

    info = Table(data, colWidths=[48*mm, 92*mm, 50*mm], rowHeights=[8.3*mm]*7)
    info.setStyle(TableStyle([
        ("GRID", (0,0), (-1,-1), .55, colors.black),
        ("VALIGN", (0,0), (-1,-1), "MIDDLE"),
        ("SPAN", (2,0), (2,6)),
        ("ALIGN", (2,0), (2,6), "CENTER"),
    ]))
    story.append(info)
    story.append(Spacer(1, 3.5*mm))

    story.append(Paragraph("EXAMINATION SCHEDULE", subtitle))
    story.append(Spacer(1, 1.5*mm))

    header = [
        Paragraph("Sr. No.", center),
        Paragraph("Date", center),
        Paragraph("Day", center),
        Paragraph("Subject / Paper", center),
        Paragraph("Time", center),
        Paragraph("Jr. Supervisor<br/>Signature", center),
    ]
    sd = [header]
    for i in range(8):
        if i < len(schedule):
            row = schedule.iloc[i]
            vals = [row.get(c, "") for c in SCHEDULE_COLS]
        else:
            vals = [i+1, "", "", "", ""]
        sd.append([
            Paragraph(esc(vals[0]), center),
            Paragraph(esc(vals[1]), normal),
            Paragraph(esc(vals[2]), normal),
            Paragraph(esc(vals[3]), normal),
            Paragraph(esc(vals[4]), center),
            ""
        ])

    stbl = Table(
        sd,
        colWidths=[15*mm, 28*mm, 28*mm, 59*mm, 28*mm, 32*mm],
        rowHeights=[8.5*mm] + [7.3*mm]*8
    )
    stbl.setStyle(TableStyle([
        ("GRID", (0,0), (-1,-1), .55, colors.black),
        ("VALIGN", (0,0), (-1,-1), "MIDDLE"),
        ("ALIGN", (0,0), (2,-1), "CENTER"),
        ("ALIGN", (4,0), (5,-1), "CENTER"),
    ]))
    story.append(stbl)
    story.append(Spacer(1, 4*mm))

    sig = Table(
        [["", "", ""],
         [Paragraph("Student Signature", center),
          Paragraph("Examination In-charge", center),
          Paragraph("Principal / Head of Institution", center)]],
        colWidths=[63.5*mm, 63.5*mm, 63.5*mm],
        rowHeights=[10*mm, 6*mm]
    )
    sig.setStyle(TableStyle([
        ("GRID", (0,0), (-1,-1), .55, colors.black),
        ("VALIGN", (0,0), (-1,-1), "MIDDLE"),
        ("ALIGN", (0,1), (-1,1), "CENTER"),
    ]))
    story.append(sig)
    story.append(Spacer(1, 3.5*mm))

    story.append(Paragraph("IMPORTANT INSTRUCTIONS", bold))
    instructions = [
        "1. The candidate must bring this Hall Ticket to every examination.",
        "2. The candidate should report to the examination centre before the scheduled time.",
        "3. Mobile phones and other prohibited electronic devices are not permitted in the examination hall.",
        "4. The candidate must follow all examination rules and instructions issued by the institution / university.",
        "5. The Hall Ticket is valid only for the examination mentioned above."
    ]
    for item in instructions:
        story.append(Paragraph(esc(item), normal))
        story.append(Spacer(1, .7*mm))

    doc.build(story)

def libreoffice_pdf(docx_path, pdf_dir):
    soffice = shutil.which("soffice") or shutil.which("libreoffice")
    if not soffice:
        return None
    try:
        subprocess.run(
            [soffice, "--headless", "--convert-to", "pdf",
             "--outdir", str(pdf_dir), str(docx_path)],
            check=True, stdout=subprocess.PIPE, stderr=subprocess.PIPE, timeout=90
        )
        p = Path(pdf_dir) / (Path(docx_path).stem + ".pdf")
        return p if p.exists() else None
    except Exception:
        return None

st.title("🎫 Automated Hall Ticket Generator")
st.caption("Corrected version — one-page Word + clean PDF output.")

with st.sidebar:
    exam_title = st.text_input("Examination title", "SEMESTER END EXAMINATION")

student_file = st.file_uploader("Upload Student Excel (.xlsx)", type=["xlsx"])
photo_files = st.file_uploader(
    "Upload photographs (optional) — filename must match Photo Filename in Excel",
    type=["jpg", "jpeg", "png"], accept_multiple_files=True
)

if student_file:
    try:
        xls = pd.ExcelFile(student_file)
        students = pd.read_excel(student_file, sheet_name="StudentData", dtype=str).fillna("")
        schedule = pd.read_excel(student_file, sheet_name="ExamSchedule", dtype=str).fillna("")

        m1 = [c for c in STUDENT_COLS if c not in students.columns]
        m2 = [c for c in SCHEDULE_COLS if c not in schedule.columns]
        if m1 or m2:
            if m1: st.error("Missing StudentData columns: " + ", ".join(m1))
            if m2: st.error("Missing ExamSchedule columns: " + ", ".join(m2))
            st.stop()

        st.success(f"{len(students)} student(s) loaded.")
        if st.button("🚀 Generate Word + PDF Hall Tickets", type="primary"):
            with tempfile.TemporaryDirectory() as td:
                td = Path(td)
                wd = td / "Word_Hall_Tickets"; pd_ = td / "PDF_Hall_Tickets"; photos = td / "photos"
                wd.mkdir(); pd_.mkdir(); photos.mkdir()

                photo_map = {}
                for f in photo_files or []:
                    p = photos / Path(f.name).name
                    p.write_bytes(f.getbuffer())
                    photo_map[f.name.lower()] = p

                docs, pdfs = [], []
                for _, s in students.iterrows():
                    stem = f"{safe_name(s['Seat / Roll No.'])}_{safe_name(s['Student Name'])}"
                    photo_name = str(s.get("Photo Filename", "")).strip()
                    photo = photo_map.get(photo_name.lower()) if photo_name else None

                    dp = wd / f"{stem}.docx"
                    pp = pd_ / f"{stem}.pdf"
                    make_docx(s, schedule, exam_title, photo, dp)
                    make_pdf(s, schedule, exam_title, photo, pp)
                    docs.append(dp); pdfs.append(pp)

                zp = td / "Hall_Tickets_WORD_PDF.zip"
                with zipfile.ZipFile(zp, "w", zipfile.ZIP_DEFLATED) as z:
                    for f in docs: z.write(f, Path("Word_Hall_Tickets") / f.name)
                    for f in pdfs: z.write(f, Path("PDF_Hall_Tickets") / f.name)

                st.success(f"Generated {len(docs)} Word file(s) and {len(pdfs)} PDF file(s).")
                st.download_button(
                    "⬇️ DOWNLOAD ALL WORD + PDF HALL TICKETS",
                    data=zp.read_bytes(),
                    file_name="Hall_Tickets_WORD_PDF.zip",
                    mime="application/zip",
                    use_container_width=True
                )
else:
    st.info("Upload the Excel workbook to begin.")
